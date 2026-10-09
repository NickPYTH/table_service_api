import logging
import os
from datetime import datetime

import openpyxl
import pytz
from apscheduler.job import Job
from apscheduler.schedulers.background import BackgroundScheduler
from apscheduler.triggers.date import DateTrigger
from django.db import connection, transaction
from django_apscheduler.jobstores import DjangoJobStore
from django_apscheduler.models import DjangoJob

from api.scheduler import scheduler as global_scheduler

logger = logging.getLogger(__name__)


TZ = pytz.timezone('Asia/Yekaterinburg')

from tables.models import Row, Column, TablePermission, Table, ColumnPermission, RowPermission



BATCH_SIZE = 5000
ROW_BATCH_SIZE = 10000


def schedule_task(func, args, job_id, run_date=None, **kwargs):
    global_scheduler.add_job(
        func=func, trigger='date', run_date=run_date or datetime.now(),
        args=args, id=job_id, jobstore='default', replace_existing=True,
        misfire_grace_time=60, max_instances=1,
    )

# Права на колонки
def subtask_column_permissions(user_id, table_id):
    column_ids = Column.objects.filter(table=table_id).values_list("id", flat=True).iterator()
    ColumnPermission.objects.bulk_create([
        ColumnPermission(column=col_id, user=user_id, table=table_id) for col_id in column_ids
    ])

# Права на строки
def subtask_row_permissions(user_id, table_id):
    row_ids = Row.objects.filter(table_id=table_id).values_list('id', flat=True).iterator()
    batch = []
    for row_id in row_ids:
        batch.append(RowPermission(row=row_id, user=user_id, table=table_id))
        if len(batch) >= ROW_BATCH_SIZE:
            RowPermission.objects.bulk_create(batch, ignore_conflicts=True)
            batch.clear()
    if batch:
        RowPermission.objects.bulk_create(batch, ignore_conflicts=True)

def safe_str(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.strftime('%d.%m.%Y')
    return str(value)


def create_column(headers, table):
    columns = []
    for order, header in enumerate(headers, start=1):
        col = Column(
            table=table,
            name=header[:299],
            order=order,
            data_type=Column.ColumnType.TEXT,
        )
        columns.append(col)
    Column.objects.bulk_create(columns)
    return columns


def create_row_and_cells(row_batch, row_batch_values, columns, table):
    if not row_batch or not row_batch_values:
        return

    with transaction.atomic():
        values_sql = ", ".join(["(%s, %s, %s, CURRENT_DATE)"] * len(row_batch))
        params = []
        for r in row_batch:
            params.extend([table.id, r.order, r.created_by_id])

        with connection.cursor() as cursor:
            cursor.execute(
                f"""
                INSERT INTO tables_row (table_id, "order", created_by_id, created_date)
                VALUES {values_sql}
                RETURNING id
                """,
                params,
            )
            row_ids = [row[0] for row in cursor.fetchall()]

        if not row_ids:
            return

        # Мультистрочный INSERT для ячеек (без RETURNING — он тут не нужен)
        cell_data = []
        for row_id, row_value in zip(row_ids, row_batch_values):
            for column, column_value in zip(columns, row_value):
                cell_data.append((table.id, row_id, column.id, safe_str(column_value)))

        if cell_data:
            cell_values_sql = ", ".join(["(%s, %s, %s, %s)"] * len(cell_data))
            cell_params = []
            for cell in cell_data:
                cell_params.extend(cell)

            with connection.cursor() as cursor:
                cursor.execute(
                    f'INSERT INTO tables_cell (table_id, "row", "column", value) VALUES {cell_values_sql}',
                    cell_params,
                )

    row_batch.clear()
    row_batch_values.clear()


def fast_import_table(file_path, table_name, user_id):
    from django.contrib.auth.models import User
    try:
        if isinstance(file_path, str):
            wb = openpyxl.load_workbook(file_path, data_only=True)
        else:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            file_path = None

        ws = wb.active
        user = User.objects.get(id=user_id)
        table = Table.objects.create(title=table_name, owner=user)
        TablePermission.objects.create(
            table=table, user=user,
            can_view=True, can_edit=True, role='owner'
        )

        row_batch = []
        row_batch_values = []
        columns = []
        for index, row_value in enumerate(ws.iter_rows(values_only=True)):
            if index == 0:
                headers = [str(header).strip() for header in row_value]
                columns = create_column(headers, table)
                if len(columns) == 0:
                    return table
            else:
                row = Row(table=table, order=index, created_by=user)
                row_batch.append(row)
                row_batch_values.append(row_value)

                if len(row_batch) >= BATCH_SIZE:
                   create_row_and_cells(row_batch, row_batch_values, columns, table)
        if row_batch:
            create_row_and_cells(row_batch, row_batch_values, columns, table)

        schedule_task(
            func=subtask_column_permissions,
            args=[user.id, table.id],
            job_id=f"col_perms_{table.id}",
            run_date=datetime.now(),
        )

        schedule_task(
            func=subtask_row_permissions,
            args=[user.id, table.id],
            job_id=f"row_perms_{table.id}",
            run_date=datetime.now(),
        )

        if file_path and os.path.exists(file_path):
            os.remove(file_path)

        return table

    except Exception as e:
        if file_path and os.path.exists(file_path):
            os.remove(file_path)
        raise


def schedule_task_in_db(func, args, job_id, run_date=None,
                        misfire_grace_time=86400, max_instances=1):

    DjangoJob.objects.filter(id=job_id).delete()

    if run_date is None:
        run_date = datetime.now(TZ)
    elif run_date.tzinfo is None:
        run_date = TZ.localize(run_date)
    trigger = DateTrigger(run_date=run_date)

    fake_scheduler = BackgroundScheduler(timezone=TZ)
    store = DjangoJobStore()
    fake_scheduler.add_jobstore(store, 'default')

    job = Job(
        scheduler=fake_scheduler,
        id=job_id,
        func=func,
        args=args,
        kwargs={},
        trigger=trigger,
        executor='default',
        name=job_id,
        misfire_grace_time=misfire_grace_time,
        coalesce=True,
        max_instances=max_instances,
    )

    job.next_run_time = run_date if run_date > datetime.now(TZ) else None

    store.add_job(job)
    logger.info(f"Задача {job_id} записана, next_run_time={job.next_run_time}")
    return job




